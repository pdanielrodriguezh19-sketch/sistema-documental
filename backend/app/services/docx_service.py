from docx import Document as DocxDocument


def extract_text_from_docx(file_path: str) -> str:
    """Extrae el texto de todos los párrafos y tablas de un DOCX."""
    doc = DocxDocument(file_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts).strip()